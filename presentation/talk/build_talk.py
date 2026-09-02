"""Build the 60-minute active-inference / comfort-zone talk on the Chalmers template.

    python presentation/talk/build_talk.py [--out presentation/talk/ai_czb_talk.pptx]

Audience: colleagues who know the datasets and the comfort-zone literature but not
this model. The deck follows the brief: the paradigm, the method with an
illustration, a worked demonstration of the published model (animated, from the
authors' own deposited output), what changes between scenarios and why, crash
causation, how the comfort-zone boundary was framed in active-inference terms,
the steps taken and what each one showed, where the project stands, and where it
goes next - with one slide on the data sources.

Every slide carries speaker notes with what to say and a time budget; the budgets
sum to about 57 minutes, leaving room for questions in the hour.

Figures: presentation/talk/figures/ (built by make_talk_figures.py and
make_event_animation.py), docs/ai_scope_figures/, docs/causation_figures/ and
figures/. Rebuild those first if any analysis has moved.

IMPORTANT for a later session: once Jonas has reviewed and animated this deck by
hand, do NOT rerun this script onto the same file - it reconstructs every slide
from the template and his edits would never exist in the new file. Copy the
reviewed deck and augment it instead (see the chalmers-slide-generation-jonas
skill, "copy and augment, never rebuild").
"""

from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
TALK_FIGS = HERE / "figures"
FIG_DIRS = [TALK_FIGS,
            REPO / "docs" / "ai_scope_figures",
            REPO / "docs" / "causation_figures",
            REPO / "docs" / "czb_figures",
            REPO / "docs" / "handbook" / "figures",
            REPO / "figures"]

TEMPLATE_POTX = REPO / "presentation" / "chalmers-tekniska-ho-gskola-sv.potx"
TEMPLATE_PPTX = REPO / "presentation" / "_chalmers-template.pptx"

# Theme colours, read from the template (see the skill's colour table)
PURPLE = RGBColor(0x47, 0x2C, 0xBE)   # accent1
LILAC = RGBColor(0x67, 0x46, 0xEB)    # accent3
MUTED = RGBColor(0x9E, 0x92, 0xE8)    # accent2
BLUE = RGBColor(0x36, 0xB7, 0xF6)     # accent5
TEAL = RGBColor(0x61, 0xE9, 0xD2)     # accent6
DEEPTEAL = RGBColor(0x1B, 0x8F, 0x7A)  # readable teal for text
PINK = RGBColor(0xD9, 0x87, 0xBA)     # accent4
DEEPPINK = RGBColor(0xB0, 0x3E, 0x82)  # readable pink for text
AMBER = RGBColor(0xC4, 0x7A, 0x14)
INK = RGBColor(0x22, 0x22, 0x22)      # dk1
BEIGE = RGBColor(0xF0, 0xED, 0xE6)    # lt2
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x5A, 0x5A, 0x5A)

FONT = "Arial"
L_HEAD = 30      # Endast rubrik
SLIDE_W, SLIDE_H = 33.87, 19.05
MARGIN = 1.30
BODY_W = 31.20
BODY_TOP = 3.95
BODY_BOT = 17.55


# ---------------------------------------------------------------------------
# template plumbing
# ---------------------------------------------------------------------------
def ensure_template() -> Path:
    """python-pptx refuses a .potx; rewrite the one content-type string."""
    if (TEMPLATE_PPTX.exists()
            and TEMPLATE_PPTX.stat().st_mtime >= TEMPLATE_POTX.stat().st_mtime):
        return TEMPLATE_PPTX
    with zipfile.ZipFile(TEMPLATE_POTX) as zin, \
            zipfile.ZipFile(TEMPLATE_PPTX, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml",
                                    b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    return TEMPLATE_PPTX


def drop_existing_slides(prs) -> None:
    id_list = prs.slides._sldIdLst
    for slide_id in list(id_list):
        prs.part.drop_rel(slide_id.rId)
        id_list.remove(slide_id)


def clean(slide) -> None:
    """Remove every unfilled placeholder so nothing ships as 'Click to add text'."""
    for shape in list(slide.placeholders):
        if shape.has_text_frame:
            if not shape.text_frame.text.strip():
                shape._element.getparent().remove(shape._element)
        else:
            try:
                _ = shape.image
            except Exception:
                shape._element.getparent().remove(shape._element)


def _plain(shape, colour) -> None:
    if colour is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False


# ---------------------------------------------------------------------------
# drawing helpers - all geometry in centimetres
# ---------------------------------------------------------------------------
def text(slide, x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.08):
    """Each run: (text, size, colour, bold[, space_before_pt])."""
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, run in enumerate(runs):
        body, size, colour, bold = run[:4]
        before = run[4] if len(run) > 4 else 0
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        if before:
            para.space_before = Pt(before)
        piece = para.add_run()
        piece.text = body
        piece.font.size = Pt(size)
        piece.font.color.rgb = colour
        piece.font.bold = bold
        piece.font.name = FONT
    return box


def bullets(items, size=14, gap=9, lead=PURPLE):
    """items: str, or (label, body) rendered as a coloured lead-in plus body."""
    runs = []
    for i, item in enumerate(items):
        before = 0 if i == 0 else gap
        if isinstance(item, tuple):
            runs.append((item[0], size, lead, True, before))
            runs.append((item[1], size, INK, False, 2))
        else:
            runs.append(("–  " + item, size, INK, False, before))
    return runs


def body(slide, items, *, x=MARGIN, y=BODY_TOP, w=BODY_W, h=None, size=14,
         gap=9, lead=PURPLE):
    h = h if h is not None else (BODY_BOT - y)
    return text(slide, x, y, w, h, bullets(items, size=size, gap=gap, lead=lead))


def panel(slide, x, y, w, h, colour=BEIGE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y),
                                 Cm(w), Cm(h))
    shp.adjustments[0] = 0.06
    _plain(shp, colour)
    return shp


def rule(slide, x, y, w, colour=PURPLE, thickness_pt=3.0):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w),
                                 Pt(thickness_pt))
    _plain(bar, colour)
    return bar


def column(slide, x, y, w, h, heading, items, *, colour=PURPLE, size=12.5,
           head_size=14.5):
    rule(slide, x, y, 1.6, colour)
    text(slide, x, y + 0.30, w, 0.9, [(heading, head_size, INK, True)])
    return text(slide, x, y + 1.35, w, h, bullets(items, size=size, gap=7,
                                                  lead=colour), spacing=1.12)


def find_figure(name: str) -> Path | None:
    for d in FIG_DIRS:
        p = d / name
        if p.exists():
            return p
    print("MISSING FIGURE:", name)
    return None


def picture_fit(slide, name, x, y, box_w, box_h, *, border=True):
    """Fit to BOTH dimensions - width-only sizing runs figures off the slide."""
    path = find_figure(name)
    if path is None:
        return None
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(Cm(box_w) / iw, Cm(box_h) / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(path), Cm(x) + (Cm(box_w) - w) // 2,
                                   Cm(y) + (Cm(box_h) - h) // 2, width=w, height=h)
    if border:
        pic.line.color.rgb = RGBColor(0xD5, 0xD1, 0xC8)
        pic.line.width = Pt(0.75)
    return pic


def table(slide, x, y, col_w, rows, *, size=12, row_h=0.78, head_colour=PURPLE):
    """Light hand-drawn table; rows[0] is the header. col_w in centimetres."""
    yy = y
    for r, row in enumerate(rows):
        head = r == 0
        if head:
            rule(slide, x, yy + row_h - 0.10, sum(col_w), head_colour, 1.6)
        xx = x
        for c, cell in enumerate(row):
            text(slide, xx, yy, col_w[c] - 0.35, row_h,
                 [(str(cell), size, head_colour if head else INK,
                   head or c == 0)], spacing=1.05)
            xx += col_w[c]
        yy += row_h + (0.14 if head else 0.0)
    return yy


def notes(slide, minutes: float, script: str) -> None:
    slide.notes_slide.notes_text_frame.text = (
        "[{:.1f} min]\n\n{}".format(minutes, script.strip()))


# ---------------------------------------------------------------------------
# slide shells
# ---------------------------------------------------------------------------
def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_HEAD])
    for shape in list(slide.placeholders):
        shape._element.getparent().remove(shape._element)
    return slide


def head(prs, title_text, *, kicker=None, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[L_HEAD])
    title = slide.shapes.title
    # Set all four geometry attributes together or the placeholder collapses.
    title.left, title.top = Cm(MARGIN), Cm(1.10)
    title.width, title.height = Cm(BODY_W), Cm(2.55)
    frame = title.text_frame
    frame.word_wrap = True
    frame.text = ""
    para = frame.paragraphs[0]
    para.line_spacing = 1.05
    run = para.add_run()
    run.text = title_text
    run.font.size = Pt(23)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = FONT
    if kicker:
        text(slide, MARGIN, 0.55, BODY_W, 0.6,
             [(kicker.upper(), 11.5, PURPLE, True)])
    if sub:
        text(slide, MARGIN, 3.35, BODY_W, 0.9, [(sub, 14.5, GREY, False)])
    clean(slide)
    return slide


def divider(prs, number, title_text, note_text, minutes=0.15):
    slide = blank(prs)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Cm(SLIDE_W), Cm(SLIDE_H))
    _plain(bg, PURPLE)
    rule(slide, 2.3, 7.7, 2.3, TEAL, 5.0)
    text(slide, 2.3, 6.8, 20, 1.0, [("PART " + str(number), 16, TEAL, True)])
    text(slide, 2.3, 8.5, 28.5, 5.0, [(title_text, 34, WHITE, True)], spacing=1.12)
    notes(slide, minutes, note_text)
    return slide


# ---------------------------------------------------------------------------
# the deck
# ---------------------------------------------------------------------------
def build(out: Path) -> None:
    prs = Presentation(str(ensure_template()))
    drop_existing_slides(prs)

    # -- 1  title ----------------------------------------------------------
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Cm(SLIDE_W), Cm(SLIDE_H))
    _plain(bg, PURPLE)
    rule(s, 2.3, 6.4, 3.2, TEAL, 5.0)
    text(s, 2.3, 7.3, 29.0, 5.4,
         [("An active-inference driver model,", 38, WHITE, True),
          ("and what it did and did not do for comfort-zone boundaries",
           38, TEAL, True, 6)], spacing=1.15)
    text(s, 2.3, 13.6, 29.0, 3.0,
         [("Jonas Bärgman  ·  Chalmers University of Technology",
           17, WHITE, False),
          ("Replicating Schumann et al. (2026, Nature Communications), extending it "
           "to crash causation,", 14.5, MUTED, False, 10),
          ("and testing it as a measurement instrument for comfort-zone boundaries",
           14.5, MUTED, False, 2)], spacing=1.15)
    notes(s, 0.8, """
Open by saying what the hour is for: this is a full account of a strand of work -
the model we adopted, what we built on it, and a pre-registered result that went
against the hypothesis we started with. The honest ending is part of the talk,
not an appendix to it.
""")

    # -- 2  roadmap --------------------------------------------------------
    s = head(prs, "Nine questions, in the order they arose",
             kicker="what the hour covers")
    col_w = 9.9
    column(s, MARGIN, BODY_TOP + 0.4, col_w, 6.0, "The idea", [
        ("1", "  Why we needed a different kind of driver model"),
        ("2", "  What active inference actually claims - and what we commit to"),
        ("3", "  The method: one loop, seven parts, one illustration"),
    ], colour=BLUE)
    column(s, MARGIN + col_w + 0.75, BODY_TOP + 0.4, col_w, 6.0, "The machinery", [
        ("4", "  One event through the model's eyes (animated, from real output)"),
        ("5", "  One driver, many worlds: what changes between scenarios"),
        ("6", "  Making the model crash for interpretable reasons"),
    ], colour=DEEPTEAL)
    column(s, MARGIN + 2 * (col_w + 0.75), BODY_TOP + 0.4, col_w, 6.0,
           "The test", [
        ("7", "  Framing the comfort-zone boundary in these terms"),
        ("8", "  What we did, and what each step showed"),
        ("9", "  Where we stand, and what we want to collect next"),
    ], colour=DEEPPINK)
    panel(s, MARGIN, 11.9, BODY_W, 4.7)
    text(s, MARGIN + 0.9, 12.45, BODY_W - 1.8, 3.8,
         [("The one-sentence version.", 16, PURPLE, True),
          ("  We adopted a published, behaviorally validated active-inference driver "
           "model, used it to build a scenario-free comfort-zone measure, and then "
           "tested that measure against human data on a design built to break it. "
           "The measure lost. What survived is a per-driver comfort-zone level that "
           "is substantially shared across scenarios - and a falsification worth "
           "publishing.", 16, INK, False)], spacing=1.25)
    notes(s, 1.2, """
Give the map, then the punchline immediately. Two reasons: nobody should spend an
hour waiting to find out that the headline claim failed, and the audience listens
to the construction differently once they know it is going to be tested rather
than advertised. Say explicitly: the negative result is pre-registered, and it is
the reason the talk is worth giving.
""")

    # =====================================================================
    divider(prs, 1, "Why a driver model, and why this one", """
Ten seconds. The section answers "what problem were we trying to solve, and what
paradigm did we borrow to solve it".
""")

    # -- 4  the problem ----------------------------------------------------
    s = head(prs, "Comfort-zone boundaries are measured one scenario and one "
                  "indicator at a time",
             kicker="the problem we started from")
    body(s, [
        ("Today's practice.", "  A minimum-TTC distribution for one conflict type, "
         "a lateral clearance for another, a headway for a third - each fitted "
         "where it is convenient, each with its own units."),
        ("Cost 1 - boundaries do not compare.", "  Nothing says whether a driver "
         "who accepts 1.2 s of headway is the same driver who accepts 0.8 m of "
         "clearance past a cyclist."),
        ("Cost 2 - the indicator is a hidden model.", "  Choosing TTC rather than "
         "gap, or PET rather than arrival-time separation, smuggles in an "
         "assumption about what the driver is regulating."),
        ("Cost 3 - a quantile describes, it does not explain.", "  A percentile of "
         "an observable says where drivers were; it does not say when or why "
         "they act."),
        ("And the end goal is a distribution, not a boundary.", "  An ADAS has to "
         "trigger at some moment. The operational proposal is a percentile of the "
         "comfort-zone boundary across drivers - which needs the boundary to be "
         "one comparable number per person."),
    ], size=13.5, gap=9)
    panel(s, MARGIN, 13.9, BODY_W, 3.3, BEIGE)
    text(s, MARGIN + 0.9, 14.40, BODY_W - 1.8, 2.4,
         [("The wish:", 14.5, PURPLE, True),
          ("  one scalar, defined the same way in every scenario, whose level set is "
           "the boundary and whose population distribution can be quoted as a "
           "percentile with an interval", 14.5, INK, False)])
    notes(s, 1.8, """
This is the framing the human-factors half of the room already agrees with, so do
not labour it. The one point to make carefully is the last: the deliverable was
never "a boundary", it was a *population distribution* of boundaries, because
that is what an ADAS trigger needs. Everything downstream - the psychometric
estimator, the percentile sensitivity, the ellipse - follows from that.

If asked about the multivariate problem: combining marginal percentiles under an
"exceed both" rule silently tightens the criterion (two 80th percentiles under
independence give 4% joint exceedance, not 20%). The QUADRARUM document works
this through and proposes an elliptical joint percentile as the principled fix.
""")

    # -- 5  the paradigm ---------------------------------------------------
    s = head(prs, "The paradigm: a driver is a prediction machine, and acting is "
                  "one of two ways to reduce prediction error",
             kicker="the underlying idea")
    body(s, [
        "The brain is not a camera followed by a calculator. It continuously guesses "
        "what its senses are about to report, and treats the mismatch - the "
        "surprise - as the thing to get rid of.",
        "There are exactly two ways to get rid of it: change your mind until it fits "
        "the world, or change the world until it fits your mind.",
    ], size=16, gap=11, h=3.2)
    y = 8.6
    cw = 15.0
    panel(s, MARGIN, y, cw, 3.4, BEIGE)
    text(s, MARGIN + 0.8, y + 0.6, cw - 1.6, 2.4,
         [("Change your mind", 17, PURPLE, True),
          ("update beliefs until they fit the evidence", 15, INK, False, 6),
          ("= perception", 15, GREY, True, 4)], spacing=1.2)
    panel(s, MARGIN + cw + 1.2, y, cw, 3.4, BEIGE)
    text(s, MARGIN + cw + 2.0, y + 0.6, cw - 1.6, 2.4,
         [("Change the world", 17, PURPLE, True),
          ("act until the world fits the beliefs", 15, INK, False, 6),
          ("= action", 15, GREY, True, 4)], spacing=1.2)
    text(s, MARGIN, y + 4.3, BODY_W, 3.4,
         [("Active inference says these are not two systems but one operation "
           "running in two directions.", 16.5, INK, False),
          ("The trick that makes it work is the preference prior: the model treats "
           "the futures the driver wants as the futures it expects. Goal-seeking "
           "becomes surprise-avoidance - one currency for both.", 16.5, PURPLE, True, 10)],
         spacing=1.25)
    notes(s, 1.5, """
Two warnings worth issuing here, because they are the commonest ways to get lost:
"surprise" is not an emotion, it is a number measuring departure from what the
model expected; and "preference" is not a choice, it is a description of the
futures a driver treats as normal, encoded so that wanting and expecting become
the same quantity.

The preference prior is the piece the whole comfort-zone argument later hangs on,
so plant it firmly now.
""")

    # -- 6  lineage --------------------------------------------------------
    s = head(prs, "Nothing here is new except the last box",
             kicker="where the idea comes from")
    picture_fit(s, "lineage_talk.png", MARGIN, 4.6, BODY_W, 6.6, border=False)
    body(s, [
        ("Predictive processing is already in our literature.", "  Great "
         "expectations (Engström, Bärgman, Nilsson, Seppelt, Markkula, "
         "Piccinini and Victor, 2018) laid out this account of driving verbally, "
         "before the present model line existed."),
        ("The way we read it,", "  the Schumann model is the computational "
         "instantiation of the account that paper gave in words."),
    ], y=12.2, size=15)
    notes(s, 1.2, """
The point of the slide for this audience: this is not an import from a foreign
field. The verbal version of the argument is a paper several people in this room
are on. What the model line adds is that it runs, and that it was benchmarked
against human response times in three conflict types - one of which was held out
of tuning entirely.
""")

    # -- 7  free energy ----------------------------------------------------
    s = head(prs, "\"Free energy\" is a model-misfit score. The word energy is a "
                  "historical accident",
             kicker="demystifying the vocabulary")
    body(s, [
        "The mathematical object is an upper bound on surprise that can be computed "
        "without knowing everything about the world. That is its entire job.",
        "The formula has the same shape as a quantity in statistical physics, so the "
        "name was borrowed. Nothing thermodynamic is meant - no heat, no metabolism.",
        "Substitute \"model-misfit score\" everywhere and you lose nothing.",
    ], size=15.5, gap=10, h=4.0)
    y = 8.4
    cw = 15.0
    panel(s, MARGIN, y, cw, 6.6, BEIGE)
    text(s, MARGIN + 0.8, y + 0.7, cw - 1.6, 5.4,
         [("Present-tense misfit", 17, PURPLE, True),
          ("how badly current beliefs fit current sensations", 14, INK, False, 8),
          ("Minimizing it is perception - the belief update.", 14, GREY, False, 8),
          ("(variational free energy)", 13, GREY, False, 8)], spacing=1.25)
    panel(s, MARGIN + cw + 1.2, y, cw, 6.6, BEIGE)
    text(s, MARGIN + cw + 2.0, y + 0.7, cw - 1.6, 5.4,
         [("Future-tense misfit", 17, PURPLE, True),
          ("how badly a candidate plan is expected to fit the preferred future",
           14, INK, False, 8),
          ("Minimizing it is action selection - choose the plan whose imagined "
           "consequences least depart from normal.", 14, GREY, False, 8),
          ("(expected free energy)", 13, GREY, False, 8)], spacing=1.25)
    text(s, MARGIN, 15.9, BODY_W, 1.4,
         [("The split between these two organizes the whole model - and, later, the "
           "whole comfort-zone argument.", 15.5, PURPLE, True)])
    notes(s, 1.2, """
Do not go near the variational derivation. If someone wants it, chapter 14 of the
handbook exists for exactly that reason.

The one thing to make stick: the second score is evaluated on *imagined* futures
under the preference distribution. That is what makes the same object serve both
"is this state acceptable" and "should I act now".
""")

    # -- 8  anchors + the claim we commit to -------------------------------
    s = head(prs, "It is a container that holds versions of the models you already "
                  "use - and we commit to only one of its three claims",
             kicker="where this sits")
    rows = [
        ["Model you know", "Where it sits inside this one"],
        ["Drift-diffusion / evidence accumulation",
         "the response-timing mechanism IS an accumulator - but the rate is the "
         "model's own computed surprise, not a fitted constant"],
        ["Looming and visual thresholds",
         "perception is optical angle and expansion, with a detection threshold; "
         "detection delay emerges rather than being fitted"],
        ["Driver risk field / safety margins",
         "the preference prior is a landscape over states - which is a risk field "
         "by another name"],
        ["Zero-risk, task-difficulty homeostasis",
         "regulation toward a comfortable region, here made explicit and computable"],
        ["Optimal control",
         "planning effort is deliberately capped, and the cost function is a "
         "probability distribution - which is what lets surprise time the response"],
    ]
    table(s, MARGIN, BODY_TOP + 0.1, [11.4, 19.8], rows, size=13, row_h=1.28)
    panel(s, MARGIN, 12.6, BODY_W, 4.6, BEIGE)
    text(s, MARGIN + 0.9, 13.05, BODY_W - 1.8, 3.8,
         [("Three claims travel under the name \"active inference\".", 15.5, PURPLE, True),
          ("(1) a process theory of the brain - neurons literally do this;  "
           "(2) a universal principle of life;  (3) an engineering framework for "
           "agents that carry uncertainty honestly and time their actions by "
           "surprise.", 14, INK, False, 8),
          ("Only (3) is needed here, and only (3) is what we claim. It is testable "
           "the ordinary way: build it, benchmark it, hold scenarios out.",
           14, DEEPTEAL, True, 8)], spacing=1.25)
    notes(s, 1.5, """
This is the slide that buys credibility with a sceptical room. Say plainly that
the framework arrives with a large and contested literature, that the critics'
central worry - a hand-built preference function plus thirteen tuned parameters
is very flexible - is fair, and that the defence is conventional science: the
intersection scenario was never used for tuning and the model still predicted
human response patterns there, and the paper ships ablations we can reproduce
from the authors' own deposit.

Nothing later in the talk depends on claims (1) or (2).
""")

    # =====================================================================
    divider(prs, 2, "How the model works", """
Ten seconds. Next four slides are the method overview: the loop, its parts, the
preference terms, and the two things that make it a psychological model rather
than a controller.
""")

    # -- 10  the loop ------------------------------------------------------
    s = head(prs, "Everything happens inside one loop, run five times a second",
             kicker="the method in one picture")
    picture_fit(s, "loop_talk.png", MARGIN, 3.9, BODY_W, 12.2, border=False)
    text(s, MARGIN, 16.4, BODY_W, 1.4,
         [("There is no emergency mode. ", 15.5, PURPLE, True),
          ("The same loop runs at every timestep of every drive; what differs "
           "between quiet following and a hard conflict is which parts of the "
           "machinery are doing the work.", 15.5, INK, False)])
    notes(s, 2.0, """
Walk the loop once, left to right, then make the two points that matter.

First: the plan is normally only *patched* - shifted one step and cheaply
re-optimized. The expensive candidate generation (about 100 candidate plans,
scored against the preferred future across the whole belief cloud, ten refinement
rounds) runs only when the accumulator demands it. That budget cap is a modelling
commitment about humans, not an implementation shortcut.

Second: the accumulator is where response *time* comes from. Not sluggish senses -
the belief cloud in the demonstration you are about to see catches the lead
braking in a single 0.2 s step.
""")

    # -- 11  the components ------------------------------------------------
    s = head(prs, "Seven components, each with an input and an output",
             kicker="the parts of the loop")
    rows = [
        ["Component", "In", "Out"],
        ["The world (generative process)", "both vehicles' controls",
         "the true state, advanced 0.2 s"],
        ["Looming perception", "true state",
         "optical angle and its rate, plus own-vehicle signals, with noise growing "
         "with distance"],
        ["Beliefs (particle filter)", "previous cloud + observation",
         "75 weighted hypotheses; the spread IS the uncertainty"],
        ["Prediction (imagined futures)", "cloud + a candidate plan",
         "a bundle of 6 s futures, sampled with a norm-compliance bias"],
        ["Preferences (the landscape)", "an imagined future",
         "how preferred that future is - six independent terms"],
        ["The planner (bounded)", "cloud + preferences",
         "the kept or replaced plan; its first step goes to the vehicle"],
        ["The surprise accumulator", "the plan's scored shortfall",
         "the re-plan trigger - and, later, our comfort-zone signal"],
    ]
    table(s, MARGIN, BODY_TOP + 0.1, [9.6, 7.4, 14.2], rows, size=12.5, row_h=1.26)
    text(s, MARGIN, 15.2, BODY_W, 2.0,
         [("Sizes, for the record: ", 14, PURPLE, True),
          ("Δt = 0.2 s, horizon 30 steps (6 s), 75 particles, 100 candidate "
           "plans with an elite fraction of 0.1 over ~10 rounds, looming threshold "
           "0.00215 s⁻¹. Thirteen parameters were hand-tuned; the rest is "
           "structure. Nothing is learned from data - there is no training set.",
           14, INK, False)])
    notes(s, 1.0, """
[CUTTABLE (1st): the loop slide carries the method on its own.]

Do not read the table. Point at three rows.

Beliefs: the cloud's spread is the model's honest uncertainty - there is no
separate confidence number. A single-best-estimate tracker cannot be of two minds;
this one routinely is, which we read as one of its most human features.

Prediction: the norms live *inside* the swarm's own motion, not as a filter
applied afterwards. Each particle runs a 32-candidate tournament weighted by norm
compliance now, one step ahead, and four seconds ahead, and takes the worse of
"now" and "later". The trust cap falls out of that min() for free - when the other
vehicle is visibly misbehaving, the shared "now" score collapses and the lottery
goes uniform, so the fan opens over everything the vehicle could physically do.

The planner: capped on purpose. An optimal planner reproduces the wrong behaviour.
""")

    # -- 12  preference terms ---------------------------------------------
    s = head(prs, "The driver's own \"normal\": six independent terms, multiplied",
             kicker="the preference prior")
    picture_fit(s, "pref_terms_talk.png", MARGIN, 4.0, BODY_W, 10.0, border=False)
    body(s, [
        ("Because they are independent, every exceedance can be blamed on a named "
         "term.", "  That is the property the comfort-zone method later exploits."),
        ("The safety term is a counterfactual, and its assumptions are the "
         "boundary's location.", "  \"If the lead braked at an assumed worst case, "
         "and I responded after an assumed reaction time, would ordinary braking "
         "still suffice?\" Never quote a boundary value without stating both."),
    ], y=14.4, size=14.5, gap=8)
    notes(s, 1.5, """
Two honest notes that a careful reader of the paper will not find in it, both
established in our method review against the released code:

The closing-rate term is one-sided in the code (the SI writes it symmetric), so it
bounds the approach *rate* and does not shape the following distance itself. And
the pedal term doubles positive accelerations and penalizes total rather than
longitudinal acceleration. Neither is in the Supplementary Information.

The standing rule in this project: no absolute boundary or headway number is ever
quoted without the assumed worst-case deceleration and the reaction-time budget
that produced it.
""")

    # -- 13  what makes it a psychological model ---------------------------
    s = head(prs, "Two currencies, one ledger - and five deliberate departures "
                  "from optimality",
             kicker="what makes it a model of a driver")
    cw = 15.0
    column(s, MARGIN, BODY_TOP + 0.3, cw, 8.4, "The two values that are added", [
        ("Pragmatic value.", "  How well an imagined future matches the preferred "
         "one: progress, comfort, safety margins."),
        ("Epistemic value.", "  How much an imagined future is expected to teach "
         "the model - looking, probing, easing off to see what the other driver does."),
        ("Same units, simply added,", "  so caution and progress trade against each "
         "other with no arbitration rule. In the published runs the pragmatic part "
         "dominates and epistemic value is inert."),
    ], colour=BLUE, size=12)
    column(s, MARGIN + cw + 1.2, BODY_TOP + 0.3, cw, 8.4,
           "Design choice → the human claim behind it", [
        ("Looming, not range sensors.", "  Drivers see angles; distant threats are "
         "genuinely harder to perceive."),
        ("A particle cloud, not one estimate.", "  Drivers entertain several "
         "readings of an ambiguous scene at once."),
        ("Norm-shaped prediction.", "  Drivers expect others to behave, and withdraw "
         "trust on evidence."),
        ("A capped planning budget.", "  Drivers satisfice."),
        ("Surprise-gated re-planning.", "  Drivers act when evidence has built "
         "up, not continuously."),
    ], colour=DEEPTEAL, size=12)
    panel(s, MARGIN, 13.3, BODY_W, 3.9, BEIGE)
    text(s, MARGIN + 0.9, 13.75, BODY_W - 1.8, 3.1,
         [("What it is not.", 15.5, DEEPPINK, True),
          ("Not learned from data. Not an optimal controller. The other vehicle is "
           "not intelligent - it follows a script and never reacts to our driver, so "
           "every published result is about unilateral avoidance. And it is not "
           "fast: one simulated timestep costs seconds to tens of seconds of CPU "
           "here, which is why our comfort-zone method deliberately avoids running "
           "the loop at all.", 14, INK, False, 8)], spacing=1.25)
    notes(s, 1.3, """
The "what it is not" panel earns more credibility than it costs, and it also sets
up two later sections: the non-reactive other vehicle is why there is no
negotiation or interaction in any of this, and the cost of the closed loop is why
the comfort-zone method is a static field evaluated on recorded kinematics.

If asked why epistemic value is set to zero: the authors' own ablation found it
inert in the longitudinal scenarios (28 matched rear-end configurations, collision
difference exactly 0.000), so our alpha = 0 matches an ablation they ran
themselves. It hints at usefulness in the oncoming scenario.
""")

    # =====================================================================
    divider(prs, 3, "One event, through the model's eyes", """
Ten seconds. The next slide is the demonstration - press play and let it run at
least twice.
""")

    # -- 15  the animation -------------------------------------------------
    s = head(prs, "A rear-end conflict, moment by moment - every number read from "
                  "the authors' own deposited output",
             kicker="demonstration")
    gif = find_figure("event_anim.gif")
    if gif is not None:
        with Image.open(gif) as im:
            iw, ih = im.size
        scale = min(Cm(BODY_W) / iw, Cm(12.9) / ih)
        s.shapes.add_picture(str(gif), Cm(MARGIN), Cm(4.0),
                             width=int(iw * scale), height=int(ih * scale))
    text(s, MARGIN, 17.2, BODY_W, 1.2,
         [("Rear-end scenario, Exp_7, seed 0: 10 m/s, 10 m gap, lead brakes at "
           "6 m/s². The accumulator is reconstructed from the deposited "
           "pragmatic-value components with the run's own λ = 10⁻⁵˙⁹⁵ "
           "and threshold 1.", 12.5, GREY, False)])
    notes(s, 3.0, """
PLAY THIS. It is an animated GIF and loops on its own in slideshow mode; let it
run through twice while you narrate.

The narration, in five beats:

t = 0.0-0.6  Steady following. Nothing happens. But note the accumulator is
already filling - about 7.6% of the threshold per step, all of it from the
collision and safety terms, because a fraction of the imagined futures over a 6 s
horizon end too close. Left alone this driver would re-plan spontaneously after
about 2.6 s of nothing happening. That is a real property of the published model
and it is not in the paper.

t = 0.8  The lead brakes. In the very next belief update the cloud has snapped to
the new reality - believed lead deceleration goes from 0.00 to -2.00 m/s^2 with
essentially zero spread. Detection is NOT the bottleneck.

t = 0.8-1.4  Knowing is not yet acting. The driver knows, and keeps its plan,
because a plan is abandoned when it stops delivering the preferred future, not
when the world changes. The per-step deposit rises from 68,000 to 197,000 to
267,000 units.

t = 1.4  The account is full. Exactly one full re-plan in the whole trial. The
winner is unambiguous: brake, hard.

t = 1.6  The brake reaches the wheels. Response time 0.8 s from the lead's onset,
of which at most one 0.2 s step was detection. It comes to rest 2.05 m behind.

If someone asks: yes, this is one seed - but response timing in this condition is
nearly deterministic, identical in 31 of 32 seeds.
""")

    # -- 16  what to take from it -----------------------------------------
    s = head(prs, "Three things to take from that event",
             kicker="reading the demonstration")
    picture_fit(s, "event_static.png", MARGIN, 4.2, 19.8, 12.6, border=False)
    cw = 10.6
    column(s, MARGIN + 20.6, 4.1, cw, 4.4, "1  Detection ≠ response", [
        "The cloud caught the braking in one step. The response came 0.6 s later, "
        "when the plan - not the world - had accumulated evidence of failure.",
        "Drivers rarely miss that something moved; what takes time is concluding "
        "that the current course of action is no longer adequate.",
    ], colour=BLUE, size=11.5)
    column(s, MARGIN + 20.6, 9.2, cw, 4.4, "2  Comfort is a zero", [
        "On the realized state the deficit inside the comfortable region is exactly "
        "zero - not merely small.",
        "So leaving the comfort zone is a defined event, not a threshold on an "
        "always-positive signal. The whole method rests on this.",
    ], colour=DEEPTEAL, size=11.5)
    column(s, MARGIN + 20.6, 13.5, cw, 4.0,
           "3  One mechanism, whole episode", [
        "Steady following, detection, response timing, braking style and the "
        "come-to-rest margin all came out of the same loop.",
        "No per-phase sub-model, no mode switch, no fitted reaction time.",
    ], colour=DEEPPINK, size=11.5)
    notes(s, 1.3, """
Point 2 is the hinge of the entire talk. Say it slowly.

Be precise about the qualification, because it is the kind of thing this audience
catches: the quantity the model itself accumulates is the shortfall expected over
noisy imagined futures, and that is NOT zero inside the zone - it grades smoothly
with the gap, from about 99,000 to 5,000 units per step across the authors' own
rear-end conditions. The field we use is the same preference function evaluated
pointwise on the realized state, which is what recorded human kinematics give us,
and that one has the exact zero. The two agree on where the boundary lies.

The authors' gap-graded expectation is, if anything, independent evidence that
their preference function encodes a comfort zone.
""")

    # =====================================================================
    divider(prs, 4, "One driver, many worlds", """
Ten seconds. What is actually different between scenarios - the question that
decides whether any of this transfers.
""")

    # -- 18  the scenario diff --------------------------------------------
    s = head(prs, "The driver barely changes. The world does",
             kicker="what differs between scenarios")
    picture_fit(s, "scenario_diff_talk.png", MARGIN, 4.0, BODY_W, 13.2, border=False)
    notes(s, 1.8, """
The evidence is a column-by-column diff of the authors' own setup tables across
all baseline runs of the three scenarios - 65 parameters per run - plus a
file-level diff of the per-scenario code.

Of those 65, everything describing the DRIVER is identical across all three:
perception noise, looming threshold, every preference weight, the planning budget,
the evidence-accumulation settings, the particle count. What changes is geometry,
initial speeds, and the other vehicle's script.

Two things worth flagging while the figure is up. decoder_true.py is byte-identical
across the three except for three docstring lines - how the world becomes
observations is scenario-independent. And the code volume is almost entirely the
other vehicle's script: oncoming's is 665 lines because it solves a small
optimal-control problem to manufacture a smooth incursion. That file contains a
function called cost_function which has nothing to do with the driver's
preferences - it is stage machinery, and reading it as psychology would be a
serious misunderstanding.

This is the model's strongest structural claim, and it is what makes the held-out
intersection test meaningful: the driver that handled the rear-end scenario was
dropped into a new world, not re-engineered for it.
""")

    # -- 19  the one driver-side change + checklist ------------------------
    s = head(prs, "One driver-side parameter changes - and it is the one that says "
                  "what kind of agent you are facing",
             kicker="the rationale, and the checklist")
    panel(s, MARGIN, BODY_TOP + 0.2, BODY_W, 4.1, BEIGE)
    text(s, MARGIN + 0.9, BODY_TOP + 0.75, BODY_W - 1.8, 3.1,
         [("w_sd_model — the steering variability the driver's internal model "
           "attributes to the OTHER vehicle.", 15.5, PURPLE, True),
          ("0.0045 in rear-end; 0.4575 in both lateral scenarios - a factor of one "
           "hundred. Not a perception setting: an assumption inside the driver's "
           "head. A lead in a queue does not steer; an oncoming or crossing vehicle "
           "might. It is the single number by which the driver was told what kind of "
           "situation it is in.", 14.5, INK, False, 8)], spacing=1.25)
    cw = 15.0
    column(s, MARGIN, 9.0, cw, 7.8, "The switching checklist  (1-3 are mechanical)", [
        ("1  Stage the world.", "  Geometry, lanes, initial states, episode length."),
        ("2  Script the other agent.", "  Trigger and manoeuvre, intensity sweepable "
         "- or, as we did for the cut-in, replay a recorded trajectory."),
        ("3  Check observability.", "  Does the new agent subtend the right angles? "
         "A pedestrian is not a truck."),
    ], colour=BLUE, size=13)
    column(s, MARGIN + cw + 1.2, 9.0, cw, 7.8,
           "4-7 are where the science is, and each needs an argument", [
        ("4  Draw the lane structure into the preferences.", "  Hand geometry; there "
         "is no map format."),
        ("5  Write the other agent's norms.", "  What does normal look like for that "
         "agent here? The most judgment-heavy step, and it directly sets how paranoid "
         "the predictions are."),
        ("6  Set the assumed variability.", "  The \"what might this thing do\" dial."),
        ("7  Re-calibrate the safety assumption,", "  and check the calibration table "
         "actually covers the speeds you will run - the failure our replication hit."),
    ], colour=DEEPPINK, size=13)
    notes(s, 1.5, """
The norms slide is the interesting one. Each scenario's norm set encodes the
specific way that scenario's threat announces itself: rear-end says only "stay in
your lane"; oncoming adds "and hold your speed", so a hard-braking oncoming vehicle
is abnormal before it crosses the line; intersection draws the junction plan and
adds "obey the light".

So writing a new scenario's norms amounts to answering: what is the earliest
observable sign, in this geometry, that the other agent has stopped being ordinary?

For our cut-in that question has an answer none of the three existing norm sets
can express - a norm that depends on the other vehicle's *manoeuvre progress*. A
vehicle in the adjacent lane is normal; a vehicle straddling the boundary is
transiently normal, because a lane change is a legal manoeuvre, but only for a
plausible duration.

Honest summary of effort: steps 1-3 are days; 4-7 are where the scientific content
lives, and skipping the argument for any of them produces a model that runs and
persuades nobody.
""")

    # =====================================================================
    divider(prs, 5, "Making the model crash, for reasons", """
Ten seconds. A complete parallel workstream - and the one place in this project
where the full closed loop actually runs.
""")

    # -- 21  crash causation: what exists, what we built --------------------
    s = head(prs, "The mechanisms are not exotic - and one of them was already in "
                  "the code, switched off",
             kicker="crash causation")
    cw = 15.0
    column(s, MARGIN, BODY_TOP + 0.3, cw, 6.4, "What the released code already has", [
        ("A complete, dormant gaze system.", "  A two-state gaze variable threaded "
         "through the whole architecture: switching probabilities, a noise multiplier "
         "of 3 when gaze is off-road, two reserved belief dimensions, a preference "
         "slot set to zero - and one hard-coded line in the planner forbidding it."),
        ("Perception-quality causation.", "  Noise scales are parameters, and looming "
         "makes perceptual difficulty state-dependent for free."),
        ("Expectation-based causation.", "  Norm trust is a looked-but-did-not-expect "
         "mechanism already."),
        ("A response-vigor dial.", "  The accumulation rate λ - blunt, but honest."),
    ], colour=BLUE, size=11.5)
    column(s, MARGIN + cw + 1.2, BODY_TOP + 0.3, cw, 8.6, "What we built on it", [
        ("Five switchable components", "  around two interchangeable response "
         "processes: off-road glances, too-close following, capped deceleration, no "
         "response, and abnormal acceleration."),
        ("Run on all 5 000 QUADRIS rear-end seed scenarios,", "  with the glance and "
         "deceleration input distributions digitized from published figures and "
         "cross-checked two independent ways."),
        ("Compared against the reference", "  with the Wu et al. binning and "
         "practical-equivalence framework."),
        ("Result:", "  the attentive active-inference driver avoids 67% of the crash "
         "population, and its conditions sit closer to the reference than the "
         "response-model control (severity θ 0.148 against 0.209)."),
    ], colour=DEEPTEAL, size=11.5)
    panel(s, MARGIN, 13.2, BODY_W, 4.0, BEIGE)
    text(s, MARGIN + 0.9, 13.65, BODY_W - 1.8, 3.2,
         [("Off-road glances are implemented as actions the driver could choose, "
           "with an evidence price attached.", 15.5, PURPLE, True),
          ("The published model was simply forbidden from choosing them. We have "
           "exercised the forced-schedule half; the model pricing its own glances "
           "through epistemic value remains the largest unexercised piece of "
           "machinery in the whole framework.", 14, INK, False, 8)], spacing=1.25)
    notes(s, 1.8, """
The dormant gaze system is the most useful single finding of reading this code.
It is not a fork - it is the code that produced the published results. The earlier
paper in the line (Engstrom et al. 2024) demonstrates exactly this machinery on
uncertainty-and-looking tasks; the collision-avoidance paper switched it off to
isolate avoidance behaviour.

On the QUADRIS caveat, state it before anyone asks: every one of those 5 000
scenarios is conditioned on having crashed under the generator's own driver. So
re-simulating answers "what does a different driver do in situations that crashed
for this one" - not "what crashes does this driver produce in traffic". Exposure
reweighting leaves an effective sample size of about 45 out of 5 000. The
between-condition contrast is robust to that; absolute rates are not, and we never
quote them.
""")

    # -- 22  the glance finding -------------------------------------------
    s = head(prs, "The gaze system gates evidence, not inference - and that is a "
                  "testable behavioral difference",
             kicker="the sharpest finding of the causation work")
    body(s, [
        ("What happened when we forced off-road glances through the code's own "
         "observation gate.", "  A driver who has already registered the lead's "
         "braking keeps responding DURING the glance, at essentially the attentive "
         "onset - even under an effectively total observation blackout."),
        ("Why.", "  The belief cloud coasts forward on its own norm-shaped "
         "prediction, and the accumulator keeps filling from remembered, "
         "extrapolated evidence. Looking away blocks new observations, not "
         "inference. The belief machinery is a short-horizon simulator, not a "
         "passive sensor buffer."),
        ("The contrast.", "  The established counterfactual-behavior model assumes "
         "the opposite: no accumulation while the eyes are off the road, and a "
         "response only 0.5 s after they return."),
    ], size=14.5, gap=11, h=8.0)
    panel(s, MARGIN, 12.4, BODY_W, 4.8, BEIGE)
    text(s, MARGIN + 0.9, 12.85, BODY_W - 1.8, 3.9,
         [("So the two architectures diverge exactly when a glance begins AFTER the "
           "conflict has been registered, and coincide when the glance covers the "
           "onset.", 16, PURPLE, True),
          ("Neither paper states this. It is visible only by running both - and it is "
           "a behavioral prediction that naturalistic glance-conditional response "
           "data could decide.", 15.5, INK, False, 8)], spacing=1.25)
    notes(s, 1.5, """
This is the one slide in the crash-causation section that is a genuine scientific
claim rather than a construction report, so give it time.

Also worth saying: the attentive active-inference onsets are later and far more
variable than the control model's fixed rule - median 1.25 s versus 0.50 s after
the same anchor - and the tier-2 closed loop confirmed that this is real rather
than a surrogate artefact.
""")

    # -- 23  the dissociation + the ROPE lesson ---------------------------
    s = head(prs, "Response timing decides how many crashes happen. The scenario "
                  "decides how hard they are",
             kicker="a result that hides inside a puzzle")
    picture_fit(s, "fig_dissociation.png", MARGIN, 4.0, 19.4, 9.0, border=False)
    column(s, MARGIN + 20.2, BODY_TOP + 0.4, 10.9, 12.4,
           "A 4.7-fold difference in crash rate...", [
        "...yet their severity quartiles agree to within 0.3 m/s. About 71% of the "
        "variation in impact speed is inherited from the scenario rather than "
        "produced by the response.",
        ("The practical consequence.", "  A driver model validated only against crash "
         "severity distributions is close to unconstrained in its response timing."),
        ("And a statistics lesson.", "  Our original equivalence criterion sat at the "
         "reference's own noise floor - a perfect model could not pass it. Three "
         "innocuous-looking choices did that: equal band weights, a bootstrap that "
         "claimed 5 000 independent scenarios where the weights carry about 950, and "
         "treating the reference as a sample rather than as the target set."),
    ], colour=DEEPPINK, size=12.5)
    notes(s, 1.5, """
[CUTTABLE (3rd): a real finding, but the section survives without it.]

The dissociation looks paradoxical - severity is produced by timing and braking,
so how can severity nearly match while timing does not? The resolution is that in
this crash population they are close to independent.

The statistics half is worth 60 seconds because it generalizes. A criterion that a
perfect model cannot pass is measuring the instrument, not the thing. The test we
now use: draw a large synthetic sample from the reference so the two are identical
by construction, and check that the criterion passes. Ours did not, at the bin
count the method's own rule prescribes.

Also: do not apply this kind of quantile-band statistic to a metric with an atom
or a lattice without fixing the bands beforehand. Our braking metric has 48% of
crashes at essentially zero braking, and two of the five band edges collapse onto
the same point, which is what produced an alarming-looking theta of 1.058 that
meant almost nothing.
""")

    # -- 24  the data ------------------------------------------------------
    s = head(prs, "Everything quoted in this talk comes from one of five sources",
             kicker="the data, in one slide")
    rows = [
        ["Source", "What it is", "What we use it for", "The caveat that binds"],
        ["The authors' OSF deposit",
         "per-run output for all three published scenarios, 32 seeds per condition, "
         "3.1 GB",
         "the replication reference; their true response-time distribution; the "
         "896 trials the dry run used",
         "it is the MODEL's behavior, not a driver's"],
        ["QUADRIS",
         "5 000 synthetic rear-end scenarios at 20 Hz, every one ending in a "
         "collision, weighted by real-world frequency",
         "the reference population for the whole crash-causation comparison",
         "every scenario is conditioned on having crashed; effective sample size "
         "~45 of 5 000 after exposure reweighting"],
        ["Digitized Bärgman et al. (2024)",
         "the SHRP2 off-road glance duration distribution and a 45-crash "
         "maximum-deceleration histogram, recovered from published figures",
         "input distributions for the glance and deceleration components",
         "the deceleration histogram is coarse - about a dozen distinct values from "
         "45 crashes"],
        ["Study 1 - clip rating and button press",
         "80 participants, 4 scenarios, 67 kinematic traces; participants say when "
         "they would intervene, and the ego never responds",
         "the boundary fit, the transfer test, and the trait analysis",
         "design order is perfectly confounded; by the button sessions people had "
         "seen every clip about four times"],
        ["Study 2 - the second cut-in study",
         "288 post-onset cells spanning a six-fold range of gap at matched TTC",
         "the decisive, pre-registered comparison of criticality axes",
         "it was built to separate gap from TTC - and it did"],
    ]
    table(s, MARGIN, BODY_TOP - 0.35, [6.6, 9.4, 7.8, 7.4], rows, size=11,
          row_h=2.12)
    notes(s, 1.5, """
You know these, so this is orientation, not exposition. One line each; the column
that matters is the last one.

The rule that applies across all five, and worth saying out loud: none of them is
a measurement of a real driver in a real conflict. Two are model output, one is a
synthetic scenario population, one is digitized from published figures, and one is
human judgment about video. The project's decisive step remains a comparison
against behavior, and nothing here substitutes for it.

Study 1's most important property for us: the ego vehicle never responds - it holds
a rigorously constant speed - so the comfort-zone field along a clip is a fixed
function of time, identical for every participant and repetition.
""")

    # =====================================================================
    divider(prs, 6, "Framing the comfort-zone boundary in these terms", """
Ten seconds. This is the part of the work that is ours rather than the authors',
and it is the part that got tested.
""")

    # -- 26  the level-set idea -------------------------------------------
    s = head(prs, "Define the boundary as a level set of one field, and the "
                  "per-scenario indicators stop being separate theories",
             kicker="the move, and the rationale")
    body(s, [
        ("The claim.", "  The preference prior already scores every driving state by "
         "how far it departs from \"how this drive is supposed to go\". Define the "
         "comfort zone as the region where that departure is zero, and the boundary "
         "as one line of constant departure."),
        ("What that buys.", "  A TTC threshold, a headway and a lateral clearance "
         "become different PROJECTIONS of the same surface - which, the way we read "
         "it, is why they never agreed across scenarios in the first place."),
        ("Why this field rather than any other score.", "  Two properties. It is "
         "exactly zero inside the comfortable region, so leaving the zone is a "
         "defined event. And it is built from the same preference function whose "
         "expected shortfall times the model's responses."),
        ("And the multivariate problem dissolves rather than being solved.", "  A "
         "per-driver scalar level has a one-dimensional population distribution by "
         "construction, so a percentile of it is well defined with no ellipse - IF "
         "the field is the right scalar."),
        ("A conceptual convergence worth recording.", "  The comfort-zone literature "
         "defines the boundary as the point where the satisficing condition is "
         "violated. A preference function is a specification of acceptable outcomes "
         "rather than an optimum to be maximized. Same idea, two vocabularies."),
    ], size=14.5, gap=11)
    notes(s, 1.8, """
Be explicit that this is our move, not the authors'. They built a collision-
avoidance model; reading their preference function as a comfort-zone field is our
reinterpretation, and it is the thing that either earns its keep or does not.

The classic constructs map without remainder, which is part of why it looked so
promising: the dread-zone boundary is where no achievable action restores the
preferred future - a physics fact, one parameter away from the comfort boundary.
Extra motives are temporary reshapings of the preference prior, with computed,
falsifiable boundary shifts. That is the next slide but one.
""")

    # -- 27  the dual role ------------------------------------------------
    s = head(prs, "The distinctive claim was never the label. It was that one "
                  "scalar does two jobs",
             kicker="what made it worth the trouble")
    picture_fit(s, "field_claim_talk.png", MARGIN, 4.0, BODY_W, 10.8, border=False)
    text(s, MARGIN, 15.2, BODY_W, 2.2,
         [("No conventional indicator has that property. ", 16, PURPLE, True),
          ("A TTC threshold can be a boundary or it can drive an accumulator, but "
           "nothing makes it the same object doing both. That is what made the "
           "framing worth the cost - and it is what made it falsifiable, because "
           "both halves can be tested separately.", 16, INK, False)], spacing=1.2)
    notes(s, 1.2, """
Say the word "falsifiable" deliberately here, and say that both halves were in
fact tested, because the next three sections are those tests.

If asked what the method actually needs to run: the field along a recorded
trajectory needs only the kinematics and the preference function. No particle
filter, no planner, no GPU - microseconds per trajectory. That design decision is
what makes it applicable to naturalistic datasets at all, and it is why the
expensive closed-loop issues never blocked the comfort-zone work.
""")

    # -- 28  boundary + extra motives -------------------------------------
    s = head(prs, "The framework predicts not only that a change of motive moves "
                  "the boundary, but by how much",
             kicker="comfort, dread, and extra motives")
    picture_fit(s, "boundary_talk.png", MARGIN, 4.0, 20.4, 11.0, border=False)
    column(s, MARGIN + 21.0, BODY_TOP + 0.4, 10.1, 12.8,
           "One parameter separates them", [
        ("Dread.", "  Where no achievable action restores the preferred future - "
         "8 m/s², physics. At 15 m/s: 0.73 s of headway."),
        ("Comfort.", "  The hardest braking a driver plans around - 4 m/s². "
         "At 15 m/s: 1.67 s, and 1.5-2.3 s across ordinary speeds, which is inside "
         "the range of observed following headways."),
        ("Hurried.", "  Reaction budget 1.0 → 0.6 s moves it to 1.27 s."),
        ("Trusting the lead.", "  Assumed worst case -6 → -3 m/s² moves it "
         "to 0.42 s."),
        ("This is the most falsifiable thing the framework offers, in our opinion.",
         "  Every number here is closed form, and every one inherits the assumed "
         "worst case and the reaction-time budget - so neither may be dropped when "
         "quoting it."),
    ], colour=DEEPTEAL, size=11.8)
    notes(s, 1.5, """
The closed form was cross-checked against the numeric level set to 0.000 m over 45
speeds. That two-independent-routes check is not ceremony: it caught a sign error
that had inflated a headway boundary from 0.7 s to a perfectly plausible-looking
3.2 s. Eyeballing would never have caught it.

The comfort boundary landing inside observed following headways is a sanity check,
not a validation, and should be presented as such.

The dread/comfort split is also where Jonas's correction applies: the dread
boundary is the deceleration drivers do not voluntarily push beyond - a behavioral
limit rather than the physical maximum - so 8 m/s^2 is an upper anchor and both
levels are fitted rather than fixed in the real analysis.
""")

    # -- 29  the decisive test --------------------------------------------
    s = head(prs, "The test that earns or loses everything: one level, fitted once, "
                  "transferring across scenarios",
             kicker="how we set out to falsify it")
    cw = 15.0
    column(s, MARGIN, BODY_TOP + 0.3, cw, 7.4, "The design", [
        "Take conflict events with response onsets in at least two scenario types.",
        "Compute the field along every trajectory - kinematics only.",
        "Fit the level on one scenario, matching first exceedance to observed onset.",
        ("Apply it unchanged to the others.", "  This step is the experiment."),
        "Score against chance, against a per-scenario refit ceiling, and against "
        "matched-complexity alternatives fitted with the same freedom.",
        "Check the decomposition blames the right term in each scenario - an "
        "attribution check the indicators cannot take.",
    ], colour=BLUE, size=12)
    column(s, MARGIN + cw + 1.2, BODY_TOP + 0.3, cw, 8.6,
           "Why either outcome is worth having", [
        ("Transfer succeeds.", "  The formulation is doing real work, and a "
         "scenario-free comfort-zone metric exists."),
        ("Transfer fails.", "  The preference function is scenario-specific - which "
         "the paper's own supplementary material concedes is possible - and we know "
         "PRECISELY how, term by term. That is itself a publishable characterization "
         "of what drivers' standards share across situations."),
        ("Per-scenario indicator thresholds cannot even express this claim.",
         "  The level-set formulation stakes itself on it. That asymmetry is the "
         "whole argument for taking the framework seriously."),
    ], colour=DEEPPINK, size=12)
    panel(s, MARGIN, 13.2, BODY_W, 4.0, BEIGE)
    text(s, MARGIN + 0.9, 13.65, BODY_W - 1.8, 3.2,
         [("What we committed to in advance, and kept.", 15.5, PURPLE, True),
          ("Models, folds and decision rules committed to the repository before each "
           "decisive run. A measured sampling-noise floor, so we knew what a perfect "
           "model would score. Matched model complexity for every comparator. And a "
           "written fallback position, to be triggered by its own stated condition "
           "rather than by taste.", 14, INK, False, 8)], spacing=1.25)
    notes(s, 1.5, """
This slide is the ethical core of the talk. Make the point that we wrote down what
would count as failure before we could see the answer, and that the fallback
position - what the project's claim would become if the field lost - was also
written in advance.

That is why the negative result that follows is a result rather than a
disappointment.
""")

    # =====================================================================
    divider(prs, 7, "What we did, and what each step showed", """
Ten seconds. The steps, in the order they happened. Two of them are pre-registered
failures.
""")

    # -- 31  the dry run --------------------------------------------------
    s = head(prs, "Step 1: the pipeline works end to end - on a model's onsets, "
                  "not a human's",
             kicker="the dry run")
    body(s, [
        ("Using the authors' own simulation output as a stand-in driver:", "  across "
         "896 rear-end trials, the full pipeline - field from kinematics, one "
         "boundary level fitted to response onsets - recovered the reference model's "
         "brake onsets with a median timing error of 0.0 s (interquartile range "
         "0.2 s, onset-matching score 0.855)."),
        ("Caveat one, stated at the time.", "  The onsets were a MODEL's. This "
         "validates the pipeline, not the psychology."),
        ("Caveat two, and it mattered later.", "  The fitted level itself was weakly "
         "identified: the rear-end field rises so steeply at the boundary that very "
         "different levels score almost alike. Boundary LOCATION robust; the level "
         "not yet disciplined."),
        ("The general lesson we now apply everywhere.", "  A fit can succeed while "
         "its parameter remains unidentified. Probe the objective around the "
         "optimum, report the range rather than the point, and treat a flat "
         "direction as a finding about the model rather than an inconvenience."),
    ], size=15.5, gap=13, h=9.0)
    panel(s, MARGIN, 13.4, BODY_W, 3.8, BEIGE)
    text(s, MARGIN + 0.9, 13.9, BODY_W - 1.8, 2.8,
         [("Scenarios where the field rises gently - lateral clearance - are exactly "
           "where the level would be pinned down.", 15.5, PURPLE, True),
          ("Which is one more reason the decisive test had to be cross-scenario.",
           15.5, INK, False, 8)], spacing=1.25)
    notes(s, 1.0, """
Keep this brisk. It is the "the machinery works" slide, and the interesting part
is the caveat rather than the number.

The identifiability point recurs throughout the project: several knobs move the
same observable. Response time is moved by the accumulator rate, the looming
threshold, perception noise, prediction noise and norm trust. Boundary location is
moved by the assumed worst-case braking AND the reaction-time budget. Fitting one
of a pair from a single observable is not possible, and pretending otherwise is how
this kind of model fools people.
""")

    # -- 32  the cut-in construction --------------------------------------
    s = head(prs, "Step 2: the released preference function could not tell a "
                  "critical cut-in from a mild one",
             kicker="the construction, and the fix")
    picture_fit(s, "cutin_before_after_talk.png", MARGIN, 3.9, BODY_W, 8.0,
                border=False)
    body(s, [
        ("The symptom.", "  Two cut-ins differing greatly in criticality produced "
         "almost identical pressure, and it moved as a step rather than a ramp."),
        ("The cause is not a bug.", "  The term asks \"if the other vehicle did its "
         "worst, could I still stop?\" - and at 30 m/s the answer is no at 10 m and "
         "no at 21 m. It is correct and saturated. What it cannot do is express "
         "degrees, and it switches on a binary lane test that a cut-in spends "
         "2.4-2.6 s straddling."),
        ("The fix, and why it is not a departure from the model.", "  Replace the "
         "binary gate by the lateral overlap fraction PREDICTED at the moment of "
         "longitudinal closure, and the saturated magnitude by the relative speed "
         "remaining at impact under maximal braking. Both are parameter-free, and "
         "both reduce exactly to the released forms in every geometry the released "
         "scenarios sustain."),
    ], y=12.2, size=12.5, gap=7)
    notes(s, 1.8, """
This is the most technically satisfying piece of work in the project and it is
worth 90 seconds, because the reasoning generalizes: when a model's term saturates,
the question is not "add a parameter" but "what continuous quantity was the binary
test standing in for".

Jonas's framing was the useful one: the physically honest quantity is the
comparison of two times - the time until lateral overlap begins against the
longitudinal time-to-collision. A collision requires the overlap to have begun by
the time the gap closes. Lane-change duration is the wrong variable, consistent
with the group's own finding that it did not matter empirically: a slow lane change
into a large gap and a fast one into a small gap can share a duration and pose
entirely different problems.

Flag also that this construction is flag-gated and defaults to off, so every
published number from the released model is untouched by it.
""")

    # -- 33  the progression ----------------------------------------------
    s = head(prs, "The steps, and where each one landed",
             kicker="the progression")
    picture_fit(s, "progression_talk.png", MARGIN, 4.0, BODY_W, 13.2, border=False)
    notes(s, 1.0, """
Use this as the map for the next four slides - point at each row as you reach it,
then move on. Do not narrate all six here; three of them get their own slide.

The colour coding is deliberate: teal means the step delivered what it promised,
pink means a pre-registered rule fired against us, amber means we discovered that
an earlier piece of evidence had been uninformative all along.
""")

    # -- 34  R.1 -----------------------------------------------------------
    s = head(prs, "Step 3: the timing half failed its pre-registered test",
             kicker="review gate R.1")
    body(s, [
        ("The rule, fixed before fitting.", "  The accumulator variant should recover "
         "most of the gap to the design regressions' cross-validated error - "
         "held-out RMSE at or below 0.11 closes it; above 0.13 is failure."),
        ("The result.", "  Three structurally-argued variants all fit WORSE in "
         "sample than a static two-parameter probit (0.169 against 0.125), and the "
         "best variant's held-out RMSE was 0.255. The rule fired."),
        ("The named cause.", "  Criticality-graded responding is already present at "
         "the first two truncation points, where at most 0.05-0.15 s of post-onset "
         "evidence can exist under any plausible motor latency. That is anticipation "
         "from repeated exposure, and no evidence-gated integrator can express it."),
    ], size=15.5, gap=13, h=8.6)
    panel(s, MARGIN, 12.8, BODY_W, 4.4, BEIGE)
    text(s, MARGIN + 0.9, 13.3, BODY_W - 1.8, 3.4,
         [("The scope of the verdict, narrowed deliberately.", 15.5, PURPLE, True),
          ("This is a FAIL for the accumulator WE specified - deficit-driven drift, "
           "non-leaky, fixed bound, on a repeated-exposure stimulus set. The standard "
           "traffic architectures (looming or TTC drift, leaky accumulation, "
           "collapsing bounds) were never tested, and the result licenses no claim "
           "about evidence accumulation as a family.", 15.5, INK, False, 8)],
         spacing=1.25)
    notes(s, 1.5, """
Two things to model for the audience here.

One: the scope narrowing is not damage control, it is accuracy. We tested one
architecture. Saying "evidence accumulation fails" would be a much bigger claim
than the data support, and it would be wrong.

Two: the failure is *interpretable*, which is worth more than a failure that is
not. The anticipation account is testable in its own right, and it is why the
project now wants naturalistic data - drivers there are not primed by seeing the
same clip four times.
""")

    # -- 35  the collinearity discovery -----------------------------------
    s = head(prs, "Step 4: the evidence we had been encouraged by could not have "
                  "decided anything",
             kicker="the collinearity discovery")
    body(s, [
        ("What we thought we had.", "  On the first cut-in study the field "
         "correlated with the human response surface at about 0.90, and out-of-sample "
         "cross-validation looked reasonable. That reads as support."),
        ("What was actually true.", "  In that study the relative speed is constant "
         "across all 18 cells. So the correlation between gap and time-to-collision "
         "is 1.0000. The longitudinal dimension has ONE degree of freedom."),
        ("Which means.", "  The field's encouraging correlation was equally "
         "consistent with a driver simply thresholding the gap. The design could "
         "never have separated them, whatever the answer had been."),
        ("What we did about it.", "  A second cut-in study that breaks the "
         "collinearity over a six-fold range of gap at matched TTC - built "
         "specifically so that the field and the simple alternatives make different "
         "predictions."),
    ], size=14.5, gap=11, h=9.0)
    panel(s, MARGIN, 13.7, BODY_W, 3.5, BEIGE)
    text(s, MARGIN + 0.9, 14.15, BODY_W - 1.8, 2.6,
         [("The transferable lesson.", 15.5, PURPLE, True),
          ("  Before believing a fit, ask what else in the design could have produced "
           "it. A correlation matrix of the candidate predictors is a five-minute "
           "check, and it would have saved us a month.", 15.5, INK, False)],
         spacing=1.25)
    notes(s, 1.5, """
This one usually gets a rueful laugh, and it should - it is the most generalizable
methodological lesson in the project, and it applies to a great deal of published
surrogate-measure work.

The point to make firmly: this is not a criticism of study 1, which was designed
for a different question. It is a criticism of reading a model comparison off a
design that cannot support one.
""")

    # -- 36  R.2 -----------------------------------------------------------
    s = head(prs, "Step 5: on the design built to separate them, the field scored "
                  "worse than chance",
             kicker="review gate R.2 - the decisive one")
    picture_fit(s, "axis_scoreboard.png", MARGIN, 4.0, 19.6, 9.0, border=False)
    column(s, MARGIN + 20.4, BODY_TOP + 0.4, 10.7, 9.6, "The comparison", [
        "288 post-onset cells. Matched three-parameter threshold models (lapse plus "
        "probit) for every candidate axis. Identical leave-one-starting-TTC-out "
        "folds. Models, folds and decision rule committed before the run.",
        ("The rule fired at a difference of +0.195.", "  Robust to excluding the "
         "20 non-modal attention-check participants, and to repairing a trace-noise "
         "handicap found after the registered run."),
        ("Within matched-TTC rows,", "  the field orders only 10 of 24 rows in the "
         "observed direction (anti-ordered in 14) where the gap orders 24 of 24 - and "
         "a numerical derivation on the preference function itself shows that "
         "inconsistency is structural, not noise."),
    ], colour=DEEPPINK, size=12.5)
    text(s, MARGIN, 13.6, 19.6, 3.6,
         [("The consequence, now in force in every project document:", 15.5, PURPLE, True),
          ("no document may describe the field as validated against human data on "
           "the longitudinal dimension. And per the gate's own advance "
           "pre-commitment, the project's headline claim was restated around what "
           "survives.", 15.5, INK, False, 6)], spacing=1.2)
    notes(s, 2.0, """
Be matter-of-fact. The field lost to a three-parameter log-gap threshold, which
approaches the measured sampling-noise floor, and it lost while scoring worse than
predicting the training mean.

Say what this does and does not mean. It rules out the KINEMATIC CONTENT of the
preference field as the criticality axis, as fitted, with the lane gate engaged.
It does not test the lateral machinery - that failed separately on the cyclist
overtake - and it does not rule out some differently constructed demand measure,
though required deceleration's poor showing makes that unpromising on this data.

If someone pushes on "isn't this just a null result": no. The alternative was
matched in complexity, the folds were grouped so the comparison is genuinely
out-of-sample, and the noise floor was measured, so we know the winning model is
close to as good as anything could be on this data.
""")

    # -- 37  the method lessons -------------------------------------------
    s = head(prs, "What the discipline bought - including one error it caught in "
                  "our own analysis",
             kicker="how we know the verdicts are real")
    cw = 15.0
    column(s, MARGIN, BODY_TOP + 0.3, cw, 7.0, "The practices", [
        ("Pre-registration inside the repository.", "  Models, folds and decision "
         "rules committed before the run. It is the reason a negative result is "
         "reportable rather than arguable."),
        ("Measure the noise floor.", "  Without it you cannot tell a good model from "
         "an easy question - or a criterion no model could pass."),
        ("Verify by two independent routes.", "  Closed form against numeric field; "
         "a cheap surrogate arbitrated against the full closed loop on 23 scenarios."),
        ("Report the residual rather than absorbing it.", "  An untuned surrogate "
         "within a stated error beats a tuned one on zero."),
    ], colour=BLUE, size=11.5)
    column(s, MARGIN + cw + 1.2, BODY_TOP + 0.3, cw, 8.6,
           "And the error the discipline caught", [
        ("The symptom.", "  Pre-onset rates were ordered by criticality in a "
         "direction the field predicted backwards - which we had written up as a "
         "limitation of the field."),
        ("The cause.", "  A covariate-window defect: the lookup began seconds before "
         "the shown clip and included the manoeuvre-onset frame, where six "
         "centimetres of lateral motion project to near-full overlap."),
        ("The fix, and what it changed.", "  Match the window to the shown clip and "
         "the pre-onset prediction goes flat. The OBSERVED gradient is real - and "
         "study 2 re-reads it as exposure-driven anticipation."),
    ], colour=DEEPTEAL, size=11.5)
    panel(s, MARGIN, 12.8, BODY_W, 4.4, BEIGE)
    text(s, MARGIN + 0.9, 13.25, BODY_W - 1.8, 3.5,
         [("Two corrections we made to our own earlier claims, for the record.",
           15.5, PURPLE, True),
          ("An earlier version of the crash-causation comparison said one condition "
           "was closer \"on every single resample\"; that analysis had held one side "
           "fixed and overstated the certainty - corrected to a posterior probability "
           "of about 0.97. And a replication discrepancy we had blamed on a "
           "calibration coverage failure turned out not to exist; the coverage "
           "failure is real and is in the authors' published runs too, but it "
           "explained nothing.", 13.5, INK, False, 8)], spacing=1.25)
    notes(s, 1.3, """
[CUTTABLE (2nd): fold the one-line version into the R.2 slide instead.]

Including our own corrections is deliberate. An audience that has watched two
pre-registered rules fire against us will trust the surviving positive result more
if they have also seen us catch and publish our own mistakes.

The C1 covariate defect is the sharpest example of why "the model predicts the
wrong sign" needs a second look before it becomes a finding. We had already written
it up as a structural limitation. It was a window bug.
""")

    # =====================================================================
    divider(prs, 8, "Where we are now", """
Ten seconds. What survived, what is switched off, and what the project honestly is.
""")

    # -- 39  what survived -------------------------------------------------
    s = head(prs, "What survived is the claim one level up - and it needs no field "
                  "at all",
             kicker="the headline result")
    picture_fit(s, "trait_shared_fraction.png", MARGIN, 4.0, 20.0, 9.6,
                border=False)
    column(s, MARGIN + 20.8, BODY_TOP + 0.4, 10.3, 10.0,
           "A driver carries one comfort-zone level", [
        "All 43 participants saw all four scenarios, so the core prediction - that a "
        "driver carries ONE level - can be tested with no field at all.",
        ("It largely holds.", "  Criticality-adjusted per-driver propensity "
         "correlates +0.50 to +0.74 across all six scenario pairs, which is 0.53 to "
         "0.78 of the split-half reliability ceiling. Mean 0.69."),
        ("The other third is scenario-specific,", "  which caps every transfer test "
         "that will ever be run on this data - so transfer should be scored against "
         "0.69 of the signal, not against perfection."),
        ("Two more things stand.", "  The boundary population is well estimated on "
         "its own scale and stable to within 1% under specification changes; and "
         "transfer to the cyclist overtake works once each scenario's response floor "
         "is freed."),
    ], colour=DEEPTEAL, size=12.5)
    notes(s, 1.8, """
This is the positive result, and it is a good one: a per-driver comfort-zone trait
that is substantially shared across a cut-in, a left turn across path, a cyclist
overtake and a truck overtake, measured with no model in the loop.

Say why the reliability ceiling matters. The raw correlations look modest until you
know that split-half reliability within each scenario is 0.93 to 0.98, so 0.69 of
the ceiling is a strong trait signal rather than a weak one.

The recommended wording for the project's claim, which we settled at the gate:
each driver carries a scalar comfort-zone threshold that is substantially shared
across scenarios; the right criticality axis is an open empirical question on which
simple scene scalars currently beat the active-inference preference field.
""")

    # -- 40  what we use and do not use -----------------------------------
    s = head(prs, "What we actually use of the framework - and what runs only in "
                  "the crash-causation work",
             kicker="the honest scope")
    picture_fit(s, "stack_talk.png", MARGIN, 4.0, BODY_W, 13.2, border=False)
    notes(s, 1.0, """
[CUTTABLE (4th): the honest-position slide repeats its content in words.]

The point of the slide: "we use active inference" was always a loose description of
the deliverable. The comfort-zone method takes a deliberately thin slice - a
calibrated preference distribution, read pointwise through one member of the
surprise family, with a true zero.

Two qualifications to state. The lane-entry weight IS a closed-form prediction of
the other vehicle's motion, which is why the top row says "in part" - what is
absent is any rollout of the ego's own policies. And for these stimuli that
omission is arguably correct rather than approximate, because participants are
asked when doing nothing stops being acceptable, which is exactly the pointwise
deficit of the no-action trajectory.

The second qualification is the one that keeps us honest with other people: the
repository's crash-causation workstream runs the entire loop on all 5 000 seeds.
So "we do not do planning" is true of the deliverable and false of the repository.
""")

    # -- 41  the honest position ------------------------------------------
    s = head(prs, "The deliverable is a measurement instrument, not an "
                  "active-inference model",
             kicker="where that leaves us")
    cw = 9.9
    column(s, MARGIN, BODY_TOP + 0.4, cw, 9.0, "Load-bearing", [
        ("Psychometrics.", "  Probit threshold plus a per-trial lapse rate - the "
         "estimator behind every fitted number."),
        ("Hierarchical Bayesian latent trait.", "  One shrunk random effect per "
         "driver; the population whose percentiles ARE the deliverable."),
        ("Held-out comparison with pre-registered rules.", "  The machinery that does "
         "the ruling - and that killed both halves of our own claim."),
        ("Comfort-zone boundary theory.", "  Defines the object being measured, and "
         "why a scalar level means anything at all."),
    ], colour=DEEPTEAL, size=11.2)
    column(s, MARGIN + cw + 0.75, BODY_TOP + 0.4, cw, 9.0,
           "Comparator, and now candidate", [
        ("Kinematic surrogates.", "  Gap, TTC, required deceleration - currently the "
         "best within-scenario criticality axis, gap leading."),
        ("Per-scenario 2D state rules.", "  How the truck and left-turn constructions "
         "are now being built."),
        ("The CZB ellipse.", "  A joint Mahalanobis percentile over per-scenario "
         "observables."),
        ("Risk-field models.", "  An independently validated instance of \"keep one "
         "scalar below a threshold\"."),
    ], colour=BLUE, size=11.2)
    column(s, MARGIN + 2 * (cw + 0.75), BODY_TOP + 0.4, cw, 9.0,
           "Falsified, and switched off", [
        ("The preference field as the criticality axis.", "  Ruled against on its own "
         "scenario, pre-registered."),
        ("Deficit-driven evidence accumulation.", "  FAIL at R.1, as specified."),
        ("EFE planning, epistemic value, belief updating.", "  Present in the "
         "repository's closed-loop agent; absent from the deliverable."),
    ], colour=DEEPPINK, size=11.2)
    panel(s, MARGIN, 13.5, BODY_W, 3.7, BEIGE)
    text(s, MARGIN + 0.9, 13.95, BODY_W - 1.8, 3.0,
         [("That is not a failed project.", 15, PURPLE, True),
          ("  A pre-registered falsification of a plausible, theoretically-motivated "
           "construct - run on a design built to discriminate, against "
           "matched-complexity alternatives, with the noise floor measured - is a "
           "result in its own right. And the trait finding is a positive one that no "
           "field was needed to establish. Active inference supplied the candidate "
           "axis, a principled zero, and the vocabulary; the axis then entered a "
           "competition and lost.", 13, INK, False)], spacing=1.2)
    notes(s, 2.0, """
Deliver the last panel as the emotional centre of the talk, and then stop talking
for a beat.

Be clear that "the deliverable is a psychometric instrument" is my framing rather
than a decision anyone signed off. A reasonable alternative framing is that this is
a comparative study of criticality axes with a psychometric estimator, which puts
the emphasis on the competition rather than on the measurement. Both are defensible
and the choice matters for how a manuscript is positioned.

One limitation to volunteer before it is asked: everything about percentiles rests
on 43 crowdsourced participants watching clips in one paradigm. That supports method
development and contrast. It does not support setting a fleet trigger.
""")

    # =====================================================================
    divider(prs, 9, "Where we go, and what we want to collect", """
Ten seconds. Three model directions and a data ask.
""")

    # -- 43  model direction + percentile ---------------------------------
    s = head(prs, "The trigger question is now a policy question as much as an "
                  "estimation question",
             kicker="where we go with the model")
    picture_fit(s, "percentile_sensitivity.png", MARGIN, 4.2, 17.4, 9.8,
                border=False)
    column(s, MARGIN + 18.2, BODY_TOP + 0.4, 12.9, 10.4,
           "Three directions, in the order we would take them", [
        ("1  Build the comparator class properly.", "  Per-scenario 2D state rules "
         "for the truck and left-turn scenarios, with both routes' pros and cons "
         "written into each construction note."),
        ("2  The CZB ellipse.", "  A joint Mahalanobis percentile over per-scenario "
         "observables, with the per-driver trait mapped onto it. Promoted from "
         "fallback to deliverable candidate."),
        ("3  Ask whether the surprise elements stand on their own.", "  Next slide - "
         "and it starts with a crux, not a construction."),
        ("What the sensitivity analysis already tells us.", "  The percentile choice "
         "and its estimation error are comparable in onset terms - about 0.27 s per "
         "5 percentile points against about 0.53 s per confidence interval. And at "
         "the population median only the most critical of the three stimuli ever "
         "crosses the boundary, so the percentile decides WHETHER milder situations "
         "trigger, not only when."),
    ], colour=BLUE, size=12.5)
    notes(s, 1.5, """
The decision this informs was stated in advance: if the trigger moves more per five
percentile points than per sampling-uncertainty band, then the bottleneck is the
percentile CHOICE - a policy question - rather than the estimator, and further
estimator refinement is deprioritized. That is roughly where we are.

Do not oversell the ellipse. It is a candidate, promoted because the one-scalar
route lost the axis competition, and its design note has not been written yet.
""")

    # -- 44  the crux ------------------------------------------------------
    s = head(prs, "Every surprise measure needs something to be surprised relative to",
             kicker="the open question")
    picture_fit(s, "reference_crux.png", MARGIN, 3.8, BODY_W, 11.4, border=False)
    text(s, MARGIN, 15.7, BODY_W, 2.2,
         [("What is falsified is the preference-field deficit as the criticality "
           "axis. ", 15.5, PURPLE, True),
          ("Not the surprise family, and not the measurement library, which is "
           "validated on its own terms and works unchanged on Gaussians, mixtures, "
           "particle sets and categoricals. But \"surprise without the field\" is not "
           "yet a specification - it is a slot with nothing in it, and naming what "
           "fills it has to come first.", 15, INK, False)], spacing=1.2)
    notes(s, 1.0, """
The two candidates both relocate the question from surprise about one's own
preferred state to surprise about the world: a learned predictive model of normal
driving for this population - which is what naturalistic data would be for - or a
per-driver predictive belief. Both are already served by the library's other
interfaces.

This is the live question, so invite disagreement on it explicitly.
""")

    # -- 45  data collection ----------------------------------------------
    s = head(prs, "What we would like to collect, and exactly what makes it usable",
             kicker="the data ask")
    cw = 9.9
    column(s, MARGIN, BODY_TOP + 0.4, cw, 8.4,
           "A  Test-track conflicts", [
        ("1  Paired-vehicle kinematics at 10 Hz or better.", "  Both vehicles on a "
         "common clock and datum; headings if turning."),
        ("2  A response-onset label per trial.", "  Brake, release or steer onset - "
         "the definition matters more than the choice."),
        ("3  A per-driver identifier across trials.", "  The whole point is "
         "per-driver levels; with 1-3 this becomes a second population."),
        ("4  The staging protocol and vehicle dimensions.", "  Including the "
         "instruction wording - the instruction confound is real."),
    ], colour=BLUE, size=11.5)
    column(s, MARGIN + cw + 0.75, BODY_TOP + 0.4, cw, 8.4,
           "B  Naturalistic driving data", [
        ("1  Measure the paradigm offset instead of assuming it.", "  Fit the level "
         "on naturalistic brake onsets and compare. The single biggest upgrade "
         "available for the absolute-trigger question."),
        ("2  Exposure denominators.", "  How often routine driving crosses a "
         "candidate trigger - the nuisance rate. No ensemble can provide it."),
        ("3  A consistency check, and transfer events in the wild.", "  Drivers "
         "should mostly live inside their own comfort zones - and here there is no "
         "instruction at all."),
    ], colour=DEEPTEAL, size=11.5)
    column(s, MARGIN + 2 * (cw + 0.75), BODY_TOP + 0.4, cw, 8.4,
           "C  What a next study must do", [
        ("Break the collinearity by design.", "  Vary relative speed independently of "
         "gap, and check the predictor correlation matrix before collecting."),
        ("Design the anticipation away, or measure it.", "  Repeated exposure "
         "produces criticality-graded responding before anything has happened; "
         "naive-exposure or catch-trial structure would remove it."),
        ("Counterbalance the order, and ask about own braking.", "  Every "
         "participant here did the rating design first, so modality and practice "
         "cannot be separated."),
    ], colour=DEEPPINK, size=11.5)
    panel(s, MARGIN, 13.4, BODY_W, 2.9, BEIGE)
    text(s, MARGIN + 0.9, 13.85, BODY_W - 1.8, 2.1,
         [("None of this needs anyone to run the model.", 15.5, PURPLE, True),
          ("  The field is computed from kinematics alone, in microseconds per "
           "trajectory. The specification written for data owners is one document, "
           "and it asks for no simulation.", 14, INK, False)], spacing=1.25)
    notes(s, 2.0, """
This is the slide to leave up during questions if the room contains anyone who
holds relevant data.

Make the last panel explicit: the ask is kinematics plus onset labels plus a driver
identifier. No model runs are required of the data owner, and the requirements
document exists and can be sent today.

On column C, be careful not to sound as though study 1 was badly designed. It was
designed for a different question and it answered that one. The point is what the
NEXT design has to do to decide the axis question, which is a question nobody had
in view when it was built.
""")

    # -- 46  closing -------------------------------------------------------
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Cm(SLIDE_W), Cm(SLIDE_H))
    _plain(bg, PURPLE)
    rule(s, 2.3, 2.9, 3.2, TEAL, 5.0)
    text(s, 2.3, 3.7, 29.2, 12.6,
         [("Four things to take away", 28, WHITE, True),
          ("A published active-inference driver model reproduces human collision-"
           "avoidance behavior from one loop, with no mode switch and no fitted "
           "reaction time - and its response timing comes from evidence "
           "accumulation, not from slow senses.", 15.5, WHITE, False, 22),
          ("One driver, many worlds: of 65 configuration parameters, every one "
           "describing the driver is shared across the three published scenarios. "
           "That is what makes transfer a meaningful question at all.",
           15.5, WHITE, False, 14),
          ("Its preference function can be read as a comfort-zone field, and that "
           "framing is falsifiable. We falsified it - twice, pre-registered - and "
           "the scope of each verdict is written down.", 15.5, TEAL, True, 14),
          ("What survives is worth having: each driver carries one comfort-zone "
           "level, about 69% shared across four scenarios, measured with no field at "
           "all. The right criticality axis is now an open competition, and simple "
           "scene scalars lead it.", 15.5, TEAL, True, 14)], spacing=1.3)
    text(s, 2.3, 17.2, 29.2, 1.6,
         [("jonas.bargman@chalmers.se   ·   the 15-chapter handbook, the method "
           "review and every analysis output are in the project repository",
           13.5, MUTED, False)])
    notes(s, 1.2, """
Land on the third and fourth bullets. The talk is not "here is a method"; it is
"here is a construction, here is how we tried to break it, here is what broke and
what did not".

Then open for questions with the data slide up if anyone in the room holds data.
""")

    prs.save(str(out))
    total = 0.0
    for slide in prs.slides:
        note = slide.notes_slide.notes_text_frame.text
        if note.startswith("["):
            total += float(note[1:note.index(" min")])
    print("wrote {}  ({} slides, notes budget {:.0f} min)".format(
        out, len(prs.slides._sldIdLst), total))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=HERE / "ai_czb_talk.pptx")
    args = ap.parse_args()
    build(args.out)


if __name__ == "__main__":
    main()
